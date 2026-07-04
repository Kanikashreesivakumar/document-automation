"""
Full runtime audit of the Trade Facility mapping context.
Prints every placeholder in trade_facility.docx and its value from the context.
"""
import sys, re
sys.path.insert(0, 'd:/Form Automation/backend')
import docx as _docx
from debug import shipments
from services.shipment_service import _shipment_to_dict
from mapping import trade_facility_context

# ─── Load real shipment ───────────────────────────────────────────────────────
s = max(shipments, key=lambda x: x.updated_at)
d = _shipment_to_dict(s)
ctx = trade_facility_context(d)

# ─── Extract all placeholders from template ───────────────────────────────────
doc = _docx.Document(r'd:\Form Automation\backend\templates\trade_facility.docx')
all_text = '\n'.join(
    [c.text for t in doc.tables for r in t.rows for c in r.cells] +
    [p.text for p in doc.paragraphs]
)
template_phs = sorted(set(re.findall(r'\{\{\s*(\w+)\s*\}\}', all_text)))

# ─── Print full context ───────────────────────────────────────────────────────
print("=" * 80)
print("TRADE FACILITY CONTEXT (runtime)")
print("=" * 80)
for k, v in sorted(ctx.items()):
    print(f"  {k}: {repr(v)}")

# ─── Audit every template placeholder ─────────────────────────────────────────
print("\n" + "=" * 80)
print("PLACEHOLDER AUDIT")
print(f"{'Placeholder':<35} {'Context Value':<35} Status")
print("-" * 80)
all_ok = True
for ph in template_phs:
    val = ctx.get(ph, '<<<MISSING_IN_CONTEXT>>>')
    if val == '<<<MISSING_IN_CONTEXT>>>':
        status = 'NOT_IN_CONTEXT'
        all_ok = False
    elif val == '':
        status = 'EMPTY'
    else:
        status = 'OK'
    print(f"  {ph:<33} {repr(str(val)):<35} {status}")

print()
if all_ok:
    print("RESULT: All template placeholders exist in the mapping context.")
else:
    print("RESULT: Some placeholders are MISSING from the context. Fix required.")

# ─── Reverse check: mapping keys not used in template ─────────────────────────
print("\n" + "=" * 80)
print("REVERSE CHECK: Context keys not used in trade_facility.docx")
print("=" * 80)
unused = sorted(set(ctx.keys()) - set(template_phs))
for k in unused:
    print(f"  {k}")
