import os
import docx
import re
import json

TEMPLATE_DIR = r'd:\Form Automation\backend\templates'
all_placeholders = {}
pattern = re.compile(r'\{\{(.*?)\}\}')

for filename in os.listdir(TEMPLATE_DIR):
    if not filename.endswith('.docx'):
        continue
    file_path = os.path.join(TEMPLATE_DIR, filename)
    doc = docx.Document(file_path)
    text = ' '.join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                text += ' ' + c.text
    matches = pattern.findall(text)
    all_placeholders[filename] = sorted(list(set([m.strip() for m in matches])))

print(json.dumps(all_placeholders, indent=2))
