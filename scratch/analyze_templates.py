import docx
from pathlib import Path
import json

def analyze_doc(filepath):
    doc = docx.Document(filepath)
    text_content = []
    
    for p in doc.paragraphs:
        if p.text.strip():
            text_content.append(f"P: {p.text.strip()}")
            
    for i, table in enumerate(doc.tables):
        for j, row in enumerate(table.rows):
            row_text = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    row_text.append(text)
            if row_text:
                text_content.append(f"T{i}R{j}: {' | '.join(row_text)}")
                
    return text_content

files = [
    r"d:\Form Automation\invoice packing list.docx",
    r"d:\Form Automation\proforma invoice.docx",
    r"d:\Form Automation\export insurance.docx",
    r"d:\Form Automation\trade facility.docx",
]

for f in files:
    print(f"\n--- {Path(f).name} ---")
    content = analyze_doc(f)
    print("\n".join(content[:20])) # print first 20 lines to get an idea
