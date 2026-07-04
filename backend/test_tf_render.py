"""
Direct end-to-end render test for trade_facility.docx.
Generates the docx, then reads it back to verify placeholders were replaced.
"""
import sys, re
sys.path.insert(0, 'd:/Form Automation/backend')
from debug import shipments
from services.shipment_service import _shipment_to_dict
from mapping import trade_facility_context
from services.docx_template_engine import fill_docx_template
import docx

s = max(shipments, key=lambda x: x.updated_at)
d = _shipment_to_dict(s)
ctx = trade_facility_context(d)

TEMPLATE = r'd:\Form Automation\backend\templates\trade_facility.docx'
OUTPUT   = r'd:\Form Automation\backend\generated\test_tf_render.docx'

print('=== CONTEXT FOR 7 FIELDS ===')
for f in ['branch_code', 'bin_number', 'starting_time', 'completion_time', 'time_taken', 'signatory_name', 'signatory_designation']:
    print(f'  {f}: {repr(ctx.get(f, "MISSING"))}')

print()
print('Rendering template...')
fill_docx_template(TEMPLATE, OUTPUT, ctx)
print(f'Rendered to: {OUTPUT}')

# Read back and check
doc = docx.Document(OUTPUT)
all_text = '\n'.join([c.text for t in doc.tables for r in t.rows for c in r.cells] + [p.text for p in doc.paragraphs])
remaining_phs = sorted(set(re.findall(r'\{\{\s*(\w+)\s*\}\}', all_text)))

print()
if remaining_phs:
    print('=== UNFILLED PLACEHOLDERS (still in output) ===')
    for ph in remaining_phs:
        print(f'  {{ {ph} }}')
else:
    print('SUCCESS: No unfilled placeholders in rendered output.')

# Show relevant cell contents
print()
print('=== RENDERED VALUES (relevant cells) ===')
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            txt = c.text.strip()
            if any(kw in txt.lower() for kw in ['branch', 'signatory', 'stuffing', 'time', 'completion', 'bin']):
                print(f'  CELL: {repr(txt[:120])}')
