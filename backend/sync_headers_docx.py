import docx
from docx.shared import Inches
import os
import zipfile

# First, extract the image from export_insurance.docx
img_path = 'templates/header_img.jpeg'
with zipfile.ZipFile('templates/export_insurance.docx', 'r') as z:
    for f in z.namelist():
        if f.startswith('word/media/'):
            with open(img_path, 'wb') as out:
                out.write(z.read(f))
            break

print("Extracted image to", img_path)

targets = ['templates/invoice.docx', 'templates/proforma_invoice.docx', 'templates/packing_list.docx']

for tgt in targets:
    doc = docx.Document(tgt)
    
    # 1. Clear the header text
    for section in doc.sections:
        for p in section.header.paragraphs:
            p.text = ""
            
    # 2. Add the image to the first paragraph of the first section's header
    header = doc.sections[0].header
    if not header.paragraphs:
        p = header.add_paragraph()
    else:
        p = header.paragraphs[0]
        
    p.alignment = 1 # Center alignment
    run = p.add_run()
    run.add_picture(img_path, width=Inches(7.5)) # Adjust width as needed
    
    # 3. Remove the hardcoded text "RASI FOODS" etc. from the body
    paragraphs_to_remove = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text
        if "RASI FOODS" in text or "MUDALAIPATTI" in text or "NAMAKKAL" in text or "rasieggs" in text or "Email" in text or "mob" in text:
            paragraphs_to_remove.append(p)
        elif i < 5 and not text.strip():
            paragraphs_to_remove.append(p)
            
    for p in paragraphs_to_remove:
        p._element.getparent().remove(p._element)
        
    doc.save(tgt)
    print(f"Updated {tgt}")

