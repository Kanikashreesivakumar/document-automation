import shutil
from pathlib import Path
from docx import Document

ROOT = Path(r"d:\Form Automation")
TEMPLATES_DIR = ROOT / "backend" / "templates"

def _replace_run_text(para, old_fragment: str, new_text: str) -> bool:
    full_text = para.text
    if old_fragment not in full_text:
        return False
        
    runs = para.runs
    positions = []
    pos = 0
    for run in runs:
        positions.append((pos, pos + len(run.text), run))
        pos += len(run.text)

    start_idx = full_text.index(old_fragment)
    end_idx = start_idx + len(old_fragment)

    first_run_idx = None
    last_run_idx = None
    for i, (rstart, rend, _) in enumerate(positions):
        if rend > start_idx and rstart < end_idx:
            if first_run_idx is None:
                first_run_idx = i
            last_run_idx = i

    if first_run_idx is None:
        return False

    first_run = positions[first_run_idx][2]
    last_run = positions[last_run_idx][2]

    before_in_first = first_run.text[:max(0, start_idx - positions[first_run_idx][0])]
    after_in_last = last_run.text[max(0, end_idx - positions[last_run_idx][0]):]

    first_run.text = before_in_first + new_text + after_in_last

    for i in range(first_run_idx + 1, last_run_idx + 1):
        positions[i][2].text = ""

    return True

def _replace_in_para(para, replacements: list):
    for old, new in replacements:
        if old in para.text:
            _replace_run_text(para, old, new)

def _replace_in_table(table, replacements: list):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_in_para(p, replacements)
            for nested in cell.tables:
                _replace_in_table(nested, replacements)

def delete_table(table):
    t = table._element
    t.getparent().remove(t)
    table._element = None

def prepare_invoice_and_packing():
    src = ROOT / "invoice packing list.docx"
    
    # 1. Prepare Invoice
    dest_inv = TEMPLATES_DIR / "invoice.docx"
    shutil.copy2(src, dest_inv)
    doc_inv = Document(str(dest_inv))
    
    # Delete table 1 (Packing list)
    if len(doc_inv.tables) > 1:
        delete_table(doc_inv.tables[1])
        
    replacements = [
        ("SHIPPING BILL NO:\tDATED:", "SHIPPING BILL NO: {{ shipping_bill_no }}\tDATED: {{ shipping_bill_date }}"),
        ("TO ORDER", "{{ consignee_name }}\n{{ consignee_address }}"),
        ("REEFER CONTAINER", "{{ means_of_transport }}"),
        ("NAMAKKAL", "{{ place_of_receipt }}"),
        ("INDIA", "{{ country_of_origin }}"),
        ("BY SEA", "{{ vessel_flight_no }}"),
        ("PORT", "{{ port_of_discharge }}"),
        ("TTNU8044030", "{{ container_no }}"),
        ("1312", "{{ total_cartons }}"),
        ("FRESH WHITE SHELL TABLE EGGS (CHICKEN). This shipment to covering under DBK scheme.", "{{ description_of_goods }}"),
        ("04072100", "{{ hsn_code }}"),
        ("50 TO 55 GMS", "{{ egg_size }}"),
        ("Amount chargeable (in words)\n US$:", "Amount chargeable (in words)\n US$: {{ amount_in_words }}"),
        ("Nett Weight :\t\nGross Weight:", "Nett Weight :\t{{ net_weight }} KGS\nGross Weight:\t{{ gross_weight }} KGS"),
    ]
    
    for p in doc_inv.paragraphs:
        _replace_in_para(p, replacements)
    for table in doc_inv.tables:
        _replace_in_table(table, replacements)
        
    doc_inv.save(str(dest_inv))
    print("Invoice done")
    
    # 2. Prepare Packing List
    dest_pl = TEMPLATES_DIR / "packing_list.docx"
    shutil.copy2(src, dest_pl)
    doc_pl = Document(str(dest_pl))
    
    # Delete table 0 (Invoice)
    if len(doc_pl.tables) > 0:
        delete_table(doc_pl.tables[0])
        
    for p in doc_pl.paragraphs:
        _replace_in_para(p, replacements)
    for table in doc_pl.tables:
        _replace_in_table(table, replacements)
        
    doc_pl.save(str(dest_pl))
    print("Packing List done")

if __name__ == "__main__":
    prepare_invoice_and_packing()
