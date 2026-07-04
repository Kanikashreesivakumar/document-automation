"""
Inject placeholders into the Container Particulars value row of trade_facility.docx.
The value row (Row 1) in Table 0 has 5 cells, all empty. We inject:
  Col 0: {{ container_no }}
  Col 1: {{ seal_no }}
  Col 2: {{ truck_no }}
  Col 3: {{ container_size }}
  Col 4: {{ total_packages }}
"""
import zipfile, shutil, re, os

SRC  = r'd:\Form Automation\backend\templates\trade_facility.docx'
DEST = r'd:\Form Automation\backend\templates\trade_facility.docx'
BAK  = r'd:\Form Automation\backend\templates\trade_facility.docx.bak'

# Back up first
if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)
    print(f"Backed up to {BAK}")

with zipfile.ZipFile(SRC, 'r') as zin:
    names = zin.namelist()
    xml = zin.read('word/document.xml').decode('utf-8')

# ── Locate the value row ────────────────────────────────────────────────────────
# The value row starts right after the header row that contains "stuffed in the container"
marker = 'stuffed in the container'
idx_marker = xml.find(marker)
assert idx_marker != -1, "Could not find header row marker"

# Find the closing </w:tr> of the header row
idx_header_tr_close = xml.find('</w:tr>', idx_marker) + len('</w:tr>')

# The value row starts at the next <w:tr
idx_val_row_start = xml.find('<w:tr ', idx_header_tr_close)
idx_val_row_end   = xml.find('</w:tr>', idx_val_row_start) + len('</w:tr>')

val_row_xml = xml[idx_val_row_start:idx_val_row_end]
print("=== ORIGINAL VALUE ROW ===")
print(val_row_xml[:400])

# ── Build the five cell run XML for each placeholder ───────────────────────────
def make_run(placeholder: str) -> str:
    return (
        f'<w:r><w:t xml:space="preserve">{placeholder}</w:t></w:r>'
    )

placeholders = [
    '{{ container_no }}',
    '{{ seal_no }}',
    '{{ truck_no }}',
    '{{ container_size }}',
    '{{ total_packages }}',
]

# ── Inject into each cell: find each <w:tc> in the value row ───────────────────
tc_pattern = re.compile(r'(<w:tc>.*?</w:tc>)', re.DOTALL)
cells = tc_pattern.findall(val_row_xml)
print(f"\nFound {len(cells)} cells in value row")

if len(cells) != 5:
    print("ERROR: Expected exactly 5 cells, got", len(cells))
    exit(1)

new_cells = []
for i, (cell_xml, ph) in enumerate(zip(cells, placeholders)):
    # Find the last <w:p ...> in the cell and insert the run before </w:p>
    # Replace the last empty paragraph with one containing our run
    # Pattern: find the <w:p ...> block and add the run inside it
    p_pattern = re.compile(r'(<w:p\b[^>]*>)(.*?)(</w:p>)', re.DOTALL)
    p_matches = list(p_pattern.finditer(cell_xml))
    if not p_matches:
        print(f"  Cell {i}: No <w:p> found, skipping")
        new_cells.append(cell_xml)
        continue

    last_p = p_matches[-1]
    p_open, p_body, p_close = last_p.group(1), last_p.group(2), last_p.group(3)
    new_p_body = p_body + make_run(ph)
    new_cell = cell_xml[:last_p.start()] + p_open + new_p_body + p_close + cell_xml[last_p.end():]
    print(f"  Cell {i}: Injected {ph}")
    new_cells.append(new_cell)

# ── Rebuild the value row ───────────────────────────────────────────────────────
# Replace cells back into the row
new_val_row = val_row_xml
# Replace each original cell with the new one in order
for orig, new in zip(cells, new_cells):
    new_val_row = new_val_row.replace(orig, new, 1)

# ── Rebuild the document XML ───────────────────────────────────────────────────
new_xml = xml[:idx_val_row_start] + new_val_row + xml[idx_val_row_end:]

# ── Write out ─────────────────────────────────────────────────────────────────
TEMP = SRC + '.tmp'
with zipfile.ZipFile(SRC, 'r') as zin:
    with zipfile.ZipFile(TEMP, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == 'word/document.xml':
                zout.writestr(item, new_xml.encode('utf-8'))
            else:
                zout.writestr(item, zin.read(item.filename))

shutil.move(TEMP, DEST)
print("\n=== TEMPLATE UPDATED ===")

# ── Verify ────────────────────────────────────────────────────────────────────
import docx, re as re2
doc = docx.Document(DEST)
t = doc.tables[0]
print(f"\nVerification - Container Particulars value row (Row 1):")
for ci, c in enumerate(t.rows[1].cells):
    txt = c.text.strip()
    print(f"  Col {ci}: {repr(txt)}")

all_phs = set(re2.findall(r'\{\{\s*(\w+)\s*\}\}', '\n'.join([c.text for r in t.rows for c in r.cells])))
print(f"\nAll placeholders in table: {all_phs}")
