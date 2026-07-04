import shutil
from pathlib import Path
from docx import Document

ROOT = Path(r"d:\Form Automation")
TEMPLATES_DIR = ROOT / "backend" / "templates"

def _replace_run_text(para, old_fragment: str, new_text: str) -> bool:
    full_text = para.text
    if old_fragment not in full_text:
        return False
        
    runs = para.runs
    positions = []
    pos = 0
    for run in runs:
        positions.append((pos, pos + len(run.text), run))
        pos += len(run.text)

    start_idx = full_text.index(old_fragment)
    end_idx = start_idx + len(old_fragment)

    first_run_idx = None
    last_run_idx = None
    for i, (rstart, rend, _) in enumerate(positions):
        if rend > start_idx and rstart < end_idx:
            if first_run_idx is None:
                first_run_idx = i
            last_run_idx = i

    if first_run_idx is None:
        return False

    first_run = positions[first_run_idx][2]
    last_run = positions[last_run_idx][2]

    before_in_first = first_run.text[:max(0, start_idx - positions[first_run_idx][0])]
    after_in_last = last_run.text[max(0, end_idx - positions[last_run_idx][0]):]

    first_run.text = before_in_first + new_text + after_in_last

    for i in range(first_run_idx + 1, last_run_idx + 1):
        positions[i][2].text = ""

    return True

def _replace_in_para(para, replacements: list):
    for old, new in replacements:
        if old in para.text:
            _replace_run_text(para, old, new)

def _replace_in_table(table, replacements: list):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_in_para(p, replacements)
                # handle nested tables
            for nested in cell.tables:
                _replace_in_table(nested, replacements)

def delete_table(doc, table_idx):
    if table_idx < len(doc.tables):
        table = doc.tables[table_idx]
        t = table._element
        t.getparent().remove(t)
        table._element = None

def prepare_export_insurance():
    src = ROOT / "export insurance.docx"
    dest = TEMPLATES_DIR / "export_insurance.docx"
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    
    replacements = [
        ("DATE:", "DATE: {{ date }}"),
        ("Date of loading                   :          ", "Date of loading                   : {{ invoice_date }}"),
        ("Container no                     :          ", "Container no                     : {{ container_no }}"),
        ("Truck                                  :        ", "Truck                                  : {{ truck_no }}"),
        ("Sum Assured                     :", "Sum Assured                     : {{ sum_assured }}"),
        ("Dollar                                 :         ", "Dollar                                 : {{ dollar_value }}"),
        ("Quantity of goods            :         ", "Quantity of goods            : {{ quantity_of_goods }} "),
        ("Port of Delivery               :         ", "Port of Delivery               : {{ port_of_delivery }} "),
        ("Risk Cover  \t \t: \tICCA ", "Risk Cover  \t \t: \t{{ risk_cover }} "),
        ("Place of Loading   \t:          Rasi Foods, ", "Place of Loading   \t: {{ place_of_loading }} "),
        ("Name of goods                  :         Fresh white shell table eggs(chicken). ", "Name of goods                  : {{ name_of_goods }} "),
        ("(SF. No: 311/5B,Minnampalli Village, Near Puthanchandai ,namakkal   pin code 637019) ", ""),
        ("Respected Sir,  \t \t", "Respected Sir,  \t \t{{ respected_sir }}"),
        ("Invoice no/Date                  :                                    DT : ", "Invoice no/Date                  : {{ invoice_no }}                                   DT : {{ invoice_date }}"),
        ("Seal Nos                            :", "Seal Nos                            : {{ seal_nos }}"),
        ("Name  and Address \t\nOf Importer", "Name  and Address \t\nOf Importer\n{{ importer_name }}\n{{ importer_address }}")
    ]
    
    for p in doc.paragraphs:
        _replace_in_para(p, replacements)
        if "Invoice no/Date" in p.text and "DT :" in p.text and "{{" not in p.text:
            p.text = 'Invoice no/Date                  : {{ invoice_no }}                                   DT : {{ invoice_date }}'
        if "Seal No" in p.text and "nos" not in p.text.lower() and ":          " in p.text:
            p.text = "Seal No's                            : {{ seal_nos }}"
        if "Name  and Address \t" in p.text and "Of Importer" in p.text:
            p.text = 'Name  and Address \t\nOf Importer: {{ importer_name }}\n{{ importer_address }}'
            
    doc.save(str(dest))
    print("Export Insurance done")

def prepare_trade_facility():
    src = ROOT / "trade facility.docx"
    dest = TEMPLATES_DIR / "trade_facility.docx"
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    
    replacements = [
        ("Shipping Bill No.                                                            :      ", "Shipping Bill No.                                                            : {{ shipping_bill_no }}"),
        ("NAME OF THE EXPORTER                                    :  RASI FOODS", "NAME OF THE EXPORTER                                    :  {{ exporter_name }}"),
        ("a. IEC NO.                                                                :  3215008319", "a. IEC NO.                                                                :  {{ iec_no }}"),
        ("GSTIN                                                                  : 33AASFR2685Q1Z8", "GSTIN                                                                  : {{ exporter_gstin }}"),
        ("Branch code                                                       : ", "Branch code                                                       : {{ branch_code }}"),
        ("BIN (PAN Based Business identification\n Number of the Exporter)                               : ", "BIN (PAN Based Business identification\n Number of the Exporter)                               : {{ bin_number }}"),
        ("4. Date of Examination                                               :  ", "4. Date of Examination                                               :  {{ date_of_examination }}"),
        ("       Starting Time                                                 : ", "       Starting Time                                                 : {{ starting_time }}"),
        ("       Completion Time                                          : ", "       Completion Time                                          : {{ completion_time }}"),
        ("       Time Taken for Stuffing                                : ", "       Time Taken for Stuffing                                : {{ time_taken }}"),
        ("Description of Cargo with quatity                     :  FRESH WHITE SHELL EGG  /1312 CARTONS", "Description of Cargo with quatity                     :  {{ description_of_cargo }}"),
        ("Country of final destination                               : ", "Country of final destination                               : {{ country_of_destination }}"),
        ("Signatory                                                               :  ", "Signatory                                                               :  {{ signatory_name }}\n{{ signatory_designation }}"),
        ("Export Invoice No.                                      :                               DT : ", "Export Invoice No.                                      : {{ invoice_no }}                              DT : {{ invoice_date }}"),
        ("Total No of Packages                                  :  1312  CARTONS", "Total No of Packages                                  :  {{ total_packages }} CARTONS"),
        ("Name & Address of the Consignee          : ", "Name & Address of the Consignee          : {{ consignee_name }}\n{{ consignee_address }}"),
        ("2.Starting Time (moving the container to CFS ): ", "2.Starting Time (moving the container to CFS ): {{ container_to_cfs_time }}"),
        ("The e-seal number is                                     , and the colour of the seal is White.", "The e-seal number is {{ e_seal_number }}, and the colour of the seal is {{ e_seal_colour }}."),
    ]
    
    for p in doc.paragraphs:
        _replace_in_para(p, replacements)
        if "Yes / No" in p.text:
            _replace_in_para(p, [("Yes / No", "{{ goods_description_verified }}")])
            
    for table in doc.tables:
        _replace_in_table(table, replacements)
        for row in table.rows:
            if "CONTAINER NO" in row.cells[0].text:
                cells = table.add_row().cells
                cells[0].text = "{{ container_no }}"
                cells[1].text = "{{ seal_no }}"
                cells[2].text = "{{ truck_no }}"
                cells[3].text = "{{ container_size }}"
                cells[4].text = "{{ packages_in_container }}"
                break
                
    doc.save(str(dest))
    print("Trade Facility done")

def prepare_proforma_invoice():
    src = ROOT / "proforma invoice.docx"
    dest = TEMPLATES_DIR / "proforma_invoice.docx"
    shutil.copy2(src, dest)
    doc = Document(str(dest))
    
    replacements = [
        # Table 0: Main table
        ("RASI FOODS\nNO. 1/219, MUDALAIPATTI, \nSALEM MAIN ROAD, NAMAKKAL -637003,\nTAMILNADU,INDIA.  Email:rasieggs@gmail.com\nGSTIN No:33AASFR2685Q1Z8", "{{ exporter_name }}\n{{ exporter_address }}\nGSTIN No:{{ exporter_gstin }}"),
        ("SHIPPING BILL NO:\tDATED:", "SHIPPING BILL NO: {{ shipping_bill_no }}\tDATED: {{ shipping_bill_date }}"),
        ("TO ORDER", "{{ consignee_name }}\n{{ consignee_address }}"),
        ("REEFER CONTAINER", "{{ means_of_transport }}"),
        ("NAMAKKAL", "{{ place_of_receipt }}"),
        ("INDIA", "{{ country_of_origin }}"),
        ("BY SEA", "{{ vessel_flight_no }}"),
        ("PORT", "{{ port_of_discharge }}"),
        # Table content
        ("TTNU8044030", "{{ container_no }}"),
        ("1312", "{{ total_cartons }}"),
        ("FRESH WHITE SHELL TABLE EGGS (CHICKEN). This shipment to covering under DBK scheme.", "{{ description_of_goods }}"),
        ("04072100", "{{ hsn_code }}"),
        ("50 TO 55 GMS", "{{ egg_size }}"),
    ]
    
    # Actually, Proforma Invoice has placeholders that need more precise handling. 
    # For now, let's just loop over all tables and paragraphs.
    for p in doc.paragraphs:
        _replace_in_para(p, replacements)
    for table in doc.tables:
        _replace_in_table(table, replacements)
        
    doc.save(str(dest))
    print("Proforma Invoice done")

def main():
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
    prepare_export_insurance()
    prepare_trade_facility()
    prepare_proforma_invoice()

if __name__ == "__main__":
    main()
